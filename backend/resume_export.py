"""Renders a structured resume dict as a single-column, ATS-friendly PDF."""

import io
from xml.sax.saxutils import escape

from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from export import _sanitize_for_pdf

_BODY = 9.5
_LEADING = 11.5
_MARGIN = 0.7 * inch
_FRAME_PADDING = 12  # reportlab Frame adds 6pt left and right internally

_NAME = ParagraphStyle(
    "Name", fontName="Times-Bold", fontSize=18, leading=22, alignment=TA_CENTER
)
_CONTACT = ParagraphStyle(
    "Contact", fontName="Times-Roman", fontSize=9, leading=11, alignment=TA_CENTER
)
_HEADING = ParagraphStyle(
    "Heading", fontName="Times-Bold", fontSize=10.5, leading=13,
    alignment=TA_LEFT, spaceBefore=8,
)
_BODY_TEXT = ParagraphStyle(
    "BodyText", fontName="Times-Roman", fontSize=_BODY, leading=_LEADING
)
_BULLET = ParagraphStyle(
    "Bullet",
    fontName="Times-Roman",
    fontSize=_BODY,
    leading=_LEADING,
    leftIndent=12,
    bulletIndent=2,
    spaceAfter=0,
)

_CONTENT_WIDTH = LETTER[0] - 2 * _MARGIN - _FRAME_PADDING


def _section_heading(title: str) -> list:
    return [
        Paragraph(escape(title.upper()), _HEADING),
        HRFlowable(width="100%", thickness=0.6, color="#000000", spaceBefore=1, spaceAfter=3),
    ]





def _joined(*parts: str) -> str:
    return " | ".join(escape(part) for part in parts if part)


def _entry(entry: dict) -> list:
    flow = [
        Paragraph(
            f"<b>{_joined(entry.get('title', ''), entry.get('title_right', ''))}</b>",
            _BODY_TEXT,
        ),
        Paragraph(
            f"<i>{_joined(entry.get('subtitle', ''), entry.get('subtitle_right', ''))}</i>",
            _BODY_TEXT,
        ),
    ]
    if entry.get("text"):
        flow.append(Paragraph(escape(entry["text"]), _BODY_TEXT))
    for bullet in entry.get("bullets", []):
        flow.append(Paragraph(escape(bullet), _BULLET, bulletText="•"))
    flow.append(Spacer(1, 6))
    return flow


def _skills(items: list[dict]) -> list:
    return [
        Paragraph(f"<b>{escape(item['label'])}:</b> {escape(item['value'])}", _BODY_TEXT)
        for item in items
    ]


def build_resume_pdf(resume: dict) -> bytes:
    resume = _sanitize_resume(resume)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        leftMargin=_MARGIN,
        rightMargin=_MARGIN,
        title=f"{resume.get('name', 'Resume')} — Resume",
    )

    story = [Paragraph(escape(resume.get("name", "")), _NAME), Spacer(1, 2)]
    contact = [escape(item) for item in resume.get("contact", []) if item]
    if contact:
        story.append(Paragraph(" | ".join(contact), _CONTACT))
    story.append(Spacer(1, 8))

    for section in resume.get("sections", []):
        story.extend(_section_heading(section["heading"]))
        kind = section.get("type", "entries")
        if kind == "text":
            story.append(Paragraph(escape(section["text"]), _BODY_TEXT))
            story.append(Spacer(1, 6))
        elif kind == "skills":
            story.extend(_skills(section.get("items", [])))
            story.append(Spacer(1, 6))
        else:
            for entry in section.get("entries", []):
                story.extend(_entry(entry))

    while story and isinstance(story[-1], Spacer):
        story.pop()
    doc.build(story)
    return buffer.getvalue()


def _sanitize_resume(value):
    if isinstance(value, str):
        return _sanitize_for_pdf(value)
    if isinstance(value, list):
        return [_sanitize_resume(item) for item in value]
    if isinstance(value, dict):
        return {key: _sanitize_resume(item) for key, item in value.items()}
    return value


def resume_to_text(resume: dict) -> str:
    lines = [resume.get("name", ""), " | ".join(resume.get("contact", []))]
    for section in resume.get("sections", []):
        lines.append(section["heading"].upper())
        if section.get("type") == "text":
            lines.append(section.get("text", ""))
        elif section.get("type") == "skills":
            lines += [f"{item['label']}: {item['value']}" for item in section.get("items", [])]
        else:
            for entry in section.get("entries", []):
                lines.append(f"{entry.get('title', '')} {entry.get('title_right', '')}".strip())
                lines.append(f"{entry.get('subtitle', '')} {entry.get('subtitle_right', '')}".strip())
                if entry.get("text"):
                    lines.append(entry["text"])
                lines += [f"- {bullet}" for bullet in entry.get("bullets", [])]
    return "\n".join(line for line in lines if line)


def build_resume_docx(resume: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    resume = _sanitize_resume(resume)
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(_BODY)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    right_edge = Inches(LETTER[0] / inch - 1.4)

    def para(text="", bold=False, italic=False, size=None, align=None, space_before=0, space_after=0):
        p = document.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return p

    def text_line(left, right, bold=False, italic=False):
        p = document.add_paragraph()
        text = " | ".join(part for part in (left, right) if part)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        return p

    def heading_rule(title):
        p = para(title.upper(), bold=True, size=10.5, space_before=8, space_after=3)
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        borders.append(bottom)
        p_pr.append(borders)

    para(resume.get("name", ""), bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    contact = [item for item in resume.get("contact", []) if item]
    if contact:
        para(" | ".join(contact), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    for section in resume.get("sections", []):
        heading_rule(section["heading"])
        kind = section.get("type", "entries")
        if kind == "text":
            para(section.get("text", ""))
        elif kind == "skills":
            for item in section.get("items", []):
                p = document.add_paragraph()
                p.add_run(f"{item['label']}: ").bold = True
                p.add_run(item["value"])
        else:
            for entry in section.get("entries", []):
                text_line(entry.get("title", ""), entry.get("title_right", ""), bold=True)
                text_line(entry.get("subtitle", ""), entry.get("subtitle_right", ""), italic=True)
                if entry.get("text"):
                    para(entry["text"])
                for bullet in entry.get("bullets", []):
                    if bullet.strip():
                        bullet_para = document.add_paragraph(bullet.strip(), style="List Bullet")
                        bullet_para.paragraph_format.left_indent = Inches(0.2)
                        bullet_para.paragraph_format.first_line_indent = Inches(-0.15)
                        bullet_para.paragraph_format.space_after = Pt(1)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
