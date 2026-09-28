"""Turns cover letter text into a downloadable PDF or DOCX."""

import io

import docx
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in text.strip().split("\n\n") if p.strip()]


# reportlab's built-in Helvetica uses WinAnsiEncoding, which lacks glyphs for
# these hyphen-family characters that LLMs commonly emit — they'd otherwise
# render as a black missing-glyph box.
_PDF_UNSUPPORTED_HYPHENS = {"‐": "-", "‑": "-", "‒": "-"}


def _sanitize_for_pdf(text: str) -> str:
    for bad, good in _PDF_UNSUPPORTED_HYPHENS.items():
        text = text.replace(bad, good)
    return text


def build_pdf(cover_letter: str) -> bytes:
    cover_letter = _sanitize_for_pdf(cover_letter)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
    )
    style = ParagraphStyle(
        "Body", fontName="Helvetica", fontSize=11, leading=16, spaceAfter=12
    )
    story = []
    for para in _paragraphs(cover_letter):
        # Preserve single line breaks (e.g. signature block) within a paragraph.
        story.append(Paragraph(para.replace("\n", "<br/>"), style))
    doc.build(story)
    return buffer.getvalue()


def build_docx(cover_letter: str) -> bytes:
    document = docx.Document()
    for para_text in _paragraphs(cover_letter):
        paragraph = document.add_paragraph()
        lines = para_text.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                paragraph.add_run().add_break()
            paragraph.add_run(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
